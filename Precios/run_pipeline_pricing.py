"""Pipeline de Pricing Tri-Proveedor (El Criollo, Posta, El Nuevo Emporio)
- Nombres de archivos estandarizados (posta.xlsx, emporio.xls)
- Bloqueo estricto de pastas y pizzas en Queso Muzzarella.
- Tapas de tarta desglosadas.
- Aceitunas descarozadas limpias.
"""

from __future__ import annotations

import argparse
import email.message
import html
import logging
import os
import re
import smtplib
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
LISTA_BUSQUEDA = BASE_DIR / "productos a buscar.txt"
CRIOLLO_CSV = BASE_DIR / "catalogo_criollo_raw.csv"

# --- NOMBRES ESTANDARIZADOS ---
POSTA_XLSX = BASE_DIR / "posta.xlsx"
EMPORIO_XLS = BASE_DIR / "emporio.xls"
OUTPUT_XLSX = BASE_DIR / "reporte_pricing_optimo.xlsx"

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s", datefmt="%H:%M:%S")
LOG = logging.getLogger("pricing_engine")

STOP_WORDS = {
    "un", "und", "unidad", "unidades", "kg", "kilo", "kilos", "gr", "g",
    "gramo", "gramos", "ml", "lt", "l", "litro", "litros", "pack", "packs",
    "caja", "cajas", "x", "por", "de", "del", "la", "el", "y", "sin", "gluten",
    "tacc", "tableta", "cajon", "balde", "of", "c/dto", "aprox", "bolsa", "tipo"
}

CRITICAL_QUALIFIERS = {
    "gochiso", "nutripanko", "naccato", "gauchito", "vaquero",
    "fontina", "pategras", "sardo", "semi", "estacionado",
    "vacalin", "repostero", "familiar", "tradicional",
    "sal", "entrefina", "gruesa", "fina", "celusal", "anclas",
    "milkaut", "paulina", "tregar", "calsa", "fenix", "garcia",
    "dolce", "dolca", "nescafe", "lucchetti", "favinco", "menoyo", "morixe", "44", "amargo",
    "tinto", "blanco", "muzzarella", "grano", "molido",
    "verdes", "negras", "descarozada"
}

FORBIDDEN_MATCHES = {
    "sal": {"aceite", "oliva", "vinagre"},
    "fontina": {"blanco", "crema", "unt", "untable", "chedd", "cheddar", "sardo", "pategras", "muzzarella"},
    "muzzarella": {"raviol", "ravioles", "sorrentino", "sorrentinos", "pasta", "pastas", "pizza", "pizzas", "empanada", "empanadas", "disco", "discos", "saltena", "canuelas"},
    "repostero": {"familiar", "tradicional", "san ignacio"},
    "familiar": {"repostero"},
    "tradicional": {"repostero"},
    "entrefina": {"gruesa", "fina"},
    "gruesa": {"entrefina", "fina"},
    "mostaza": {"dijon", "arytza", "savora", "aderezo", "dips", "salsa", "pasta"},
    "avellana": {"syrup", "sirope", "jarabe", "pehuenia", "salsa", "pasta", "bano", "relleno", "lheritier", "licor", "aroma", "esencia", "topping"},
    "avellanas": {"syrup", "sirope", "jarabe", "pehuenia", "salsa", "pasta", "bano", "relleno", "lheritier", "licor", "aroma", "esencia", "topping"},
    "amargo": {"alcalino", "dulce", "leche"},
    "crema": {"39", "39%", "repostera", "vegetal", "chantilly"},
    "aceite": {"alto oleico", "oleico", "oliva", "maiz", "soja"},
    "ricota": {"raviol", "ravioles", "sorrentino", "sorrentinos", "tapa", "tapas"},
    "vino": {"rosado"},
    "tinto": {"blanco", "rosado"},
    "blanco": {"tinto", "rosado"},
    "grano": {"mol", "molido", "molida", "licor", "polvo"},
    "molido": {"grano", "entero", "semilla", "semillas"},
    "verdes": {"negras", "negra", "fileteada", "fileteadas", "rellena", "rellenas", "morron"},
    "negras": {"verdes", "verde", "fileteada", "fileteadas", "rellena", "rellenas", "morron"},
    "descarozada": {"fileteada", "fileteadas", "rellena", "rellenas", "morron", "pasta"},
    "azucar": {"muñeca", "muneca"},
}

GENERIC_MULTI_OPTION_PATTERNS = [
    "harina 000",
    "harina 0000",
    "azucar",
    "azucar negra",
    "azucar impalpable",
    "pure tomate",
    "extracto tomate",
    "avellanas",
    "aceite girasol",
    "cacao amargo",
    "coco rallado",
    "pan rallado",
    "queso pategras",
    "queso muzzarella",
    "panko",
    "vino blanco",
    "vino tinto",
    "fideos",
    "miel",
    "tapa tarta",
    "aceitunas negras descarozada",
    "aceitunas verdes descarozada",
]

CANONICAL_TARGET_MAP = {
    "queso sardo semi estacionado": "Queso Sardo Semi Gauchito",
    "sardo semi estacionado": "Queso Sardo Semi Gauchito",
    "queso sardo semi": "Queso Sardo Semi Gauchito",
    "sardo semi": "Queso Sardo Semi Gauchito",
    "queso muzzarella gauchito": "Queso Muzzarella",
    "queso mozzarella gauchito": "Queso Muzzarella",
    "queso mozzarella": "Queso Muzzarella",
    "muzzarella": "Queso Muzzarella",
    "tapa pascualina": "Tapa de tarta",
    "tapa tarta": "Tapa de tarta",
    "tapas tarta": "Tapa de tarta",
    "aceitunas verdes desc": "Aceitunas verdes descarozadas",
    "aceitunas negras desc": "Aceitunas negras descarozadas",
    "mostaza grano": "Semillas de mostaza",
    "mostaza en grano": "Semillas de mostaza",
    "semillas de mostaza": "Semillas de mostaza",
    "curcuma molida": "Cúrcuma en polvo",
    "curcuma en polvo": "Cúrcuma en polvo",
    "canela molida": "Canela en polvo",
    "canela en polvo": "Canela en polvo",
    "comino molido": "Comino en polvo",
    "comino en polvo": "Comino en polvo",
    "coriandro grano": "Semillas de coriandro",
    "semillas de coriandro": "Semillas de coriandro",
    "levadura fresca calsalevadura": "Levadura fresca Calsa",
}


def normalize_text(value: Any) -> str:
    text = "" if value is None or pd.isna(value) else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    text = re.sub(r"(\d+(?:[.,]\d+)?)\s*(kg|kgs|kilo|kilos|gr|grs|g|lt|lts|l|litro|litros|ml|cc)\b", r"\1 \2", text)
    
    abbrevs = {
        r"\bddl\b": "dulce leche",
        r"\brep\b": "repostero",
        r"\bfam\b": "familiar",
        r"\brall\b": "rallado",
        r"\bcrem\b": "crema",
        r"\blech\b": "leche",
        r"\bente\b": "entera",
        r"\bpel\b": "pelada",
        r"\btto\b": "tinto",
        r"\bbco\b": "blanco",
        r"\bgochis\b": "gochiso",
        r"\bgochiso\b": "gochiso",
        r"\bestac\b": "estacionado",
        r"\bs/pint\b": "sin pintar",
        r"\bpint\b": "pintado",
        r"\b44%\b": "44",
        r"\b39%\b": "39",
        r"\ben polvo\b": "molido",
        r"\bpolvo\b": "molido",
        r"\bmolida\b": "molido",
        r"\bsemillas\b": "grano",
        r"\bsemilla\b": "grano",
        r"\be/fn\b": "entrefina",
        r"\be/f\b": "entrefina",
        r"\bluchetti\b": "lucchetti",
        r"\bvirgen\b": "fresca",
        r"\bprensada\b": "fresca",
        r"\bmuzz\b": "muzzarella",
        r"\bmozz\b": "muzzarella",
        r"\bmozzar\b": "muzzarella",
        r"\bmozzarella\b": "muzzarella",
        r"\bmuza\b": "muzzarella",
        r"\bpasc\b": "tarta",
        r"\bhorno\b": "tarta",
        r"\bpascualina\b": "tarta",
        r"\btartas\b": "tarta",
        r"\btapas\b": "tapa",
        r"\bred\b": "redonda",
        r"\bcil\b": "cilindro",
        r"\bdesc\b": "descarozada",
        r"\bdescaroz\b": "descarozada",
        r"\bdescarozado\b": "descarozada",
        r"\bdescarozados\b": "descarozada",
        r"\bdescarozadas\b": "descarozada",
        r"\bs/carozo\b": "descarozada",
        r"\bs/carozos\b": "descarozada",
        r"\bs/caroz\b": "descarozada",
        r"\bs/car\b": "descarozada",
        r"\bs/c\b": "descarozada",
    }
    for k, v in abbrevs.items():
        text = re.sub(k, v, text)
        
    text = re.sub(r"[^a-z0-9]+", " ", text)
    tokens = [t for t in text.split() if t not in STOP_WORDS]
    return " ".join(tokens)


def parse_price(val: Any) -> float | None:
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (int, float)):
        v = float(val)
        return v if v > 0 else None

    s = str(val).strip().replace("$", "").replace(" ", "")
    m = re.search(r"[\d.,]+", s)
    if not m:
        return None
    raw = m.group(0)

    if "," in raw and "." in raw:
        if raw.rfind(".") > raw.rfind(","):
            raw = raw.replace(",", "")
        else:
            raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        parts = raw.split(",")
        raw = raw.replace(",", ".") if len(parts[-1]) <= 2 else raw.replace(",", "")
    elif "." in raw:
        parts = raw.split(".")
        if len(parts[-1]) == 3 and len(parts) == 2:
            raw = raw.replace(".", "")

    try:
        res = float(raw)
        return res if res > 0 else None
    except ValueError:
        return None


def is_unit_only_item(norm_text: str) -> bool:
    tokens = set(norm_text.split())
    if {"pure", "tomate"}.issubset(tokens):
        return True
    if {"extracto", "tomate"}.issubset(tokens):
        return True
    if any(k in tokens for k in {"fideo", "fideos", "tirabuzon"}):
        return True
    if "tarta" in tokens and "tapa" in tokens:
        return True
    return False


def parse_presentation_by_unit(text: Any, unit_hint: str = "", price_nominal: float | None = None) -> tuple[float, str, str]:
    source = "" if text is None or pd.isna(text) else str(text)
    compact = re.sub(r"\s+", " ", source.strip().lower())
    norm = normalize_text(compact)
    u_hint = str(unit_hint).strip().upper()

    if is_unit_only_item(norm):
        pack_match = re.search(r"(?:x\s*|caja\s*(?:x\s*)?)(\d+)\s*(?:un|und|u\b)", compact)
        if pack_match:
            pack_units = float(pack_match.group(1))
            if price_nominal and price_nominal > 3000 and pack_units > 1:
                return pack_units, "UN", source
        return 1.0, "UN", source

    if u_hint in {"KG", "X KG", "KGS", "KILO", "KILOS"}:
        return 1.0, "KG", source
    if u_hint in {"LT", "L", "LTS", "LITRO", "LITROS"}:
        return 1.0, "LT", source

    number = r"\d+(?:[.,]\d+)?"
    mass_volume = re.search(rf"({number})\s*(kgs?|kilos?|grs?\.?|gramos?|g|lt|lts?|litros?|l|ml|cc)\b", compact)

    if mass_volume:
        amount = float(mass_volume.group(1).replace(",", "."))
        raw_unit = mass_volume.group(2).rstrip(".")
        if raw_unit in {"kg", "kgs", "kilo", "kilos"}:
            base_unit = "KG"
        elif raw_unit in {"gr", "grs", "gramo", "gramos", "g"}:
            amount /= 1000.0
            base_unit = "KG"
        elif raw_unit in {"lt", "lts", "litro", "litros", "l"}:
            base_unit = "LT"
        elif raw_unit in {"ml", "cc"}:
            amount /= 1000.0
            base_unit = "LT"
        else:
            base_unit = "UN"

        return (amount if amount > 0 else 1.0), base_unit, source

    return 1.0, "UN", source


def find_all_safe_matches(query: str, choices: list[str]) -> list[str]:
    query_tokens = set(query.split())

    filtered_choices = choices
    for key, forbidden in FORBIDDEN_MATCHES.items():
        if key in query_tokens:
            filtered_choices = [c for c in filtered_choices if not any(fb in c.split() for fb in forbidden)]

    if "estacionado" in query_tokens and "semi" not in query_tokens:
        filtered_choices = [c for c in filtered_choices if "semi" not in c.split()]

    required_must_have = query_tokens & CRITICAL_QUALIFIERS

    if "semi" in query_tokens:
        required_must_have.discard("estacionado")

    if "muzzarella" in query_tokens:
        required_must_have.discard("gauchito")

    if required_must_have:
        filtered_choices = [c for c in filtered_choices if required_must_have.issubset(set(c.split()))]

    if not filtered_choices:
        return []

    def stem(token: str) -> str:
        return token.rstrip("s")

    ignore_in_scoring = {"queso", "de", "en", "el", "la"}
    numeric_terms = {token for token in query_tokens if token.isdigit()}
    key_terms = {stem(token) for token in query_tokens if len(token) >= 3 and not token.isdigit() and token not in ignore_in_scoring}

    required_overlap = min(len(key_terms), 2) if key_terms else 1
    semantic_choices = []
    
    for choice in filtered_choices:
        c_tokens = set(choice.split())
        overlap = sum(1 for term in key_terms if any(stem(ct) == term for ct in c_tokens))
        if overlap >= required_overlap and numeric_terms.issubset(c_tokens):
            semantic_choices.append(choice)

    return semantic_choices


def is_generic_item(item_text: str) -> bool:
    norm = normalize_text(item_text)
    brand_keywords = {
        "milkaut", "paulina", "celusal", "anclas", "tregar", "calsa",
        "fenix", "garcia", "emeth", "lucchetti", "favinco", "menoyo",
        "dolce", "dolca", "vacalin", "gochiso"
    }
    if any(b in norm.split() for b in brand_keywords):
        return False
    norm_tokens = set(norm.split())
    for pat in GENERIC_MULTI_OPTION_PATTERNS:
        pat_tokens = set(normalize_text(pat).split())
        if pat_tokens.issubset(norm_tokens):
            return True
    return False


def extract_keywords_from_targets(targets: list[str]) -> list[str]:
    kw_set = {
        "pategras", "fontina", "sardo", "rallado", "morixe", "vacalin", "lucchetti",
        "tirabuzon", "anclas", "celusal", "curcuma", "comino", "canela", "coco",
        "mostaza", "coriandro", "tomillo", "oregano", "hongos", "pino", "dolce",
        "dolca", "nescafe", "favinco", "menoyo", "avellana", "avellanas", "fenix",
        "garcia", "tregar", "milkaut", "paulina", "girasol", "cacao", "pure", "extracto",
        "sal", "panko", "gochiso", "gauchito", "muzzarella", "muzz", "tarta", "pascualina",
        "vino", "nativo", "miel", "aceitunas", "castell", "reygal"
    }

    skip = {"de", "en", "el", "la", "los", "las", "x", "kg", "gr", "un", "para", "sin", "gluten", "polvo", "seco", "seca", "entero", "entera", "molido", "molida"}

    for t in targets:
        clean = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode().lower()
        clean = re.sub(r"[^a-z0-9]+", " ", clean)
        tokens = [tok for tok in clean.split() if tok not in skip and not tok.isdigit() and len(tok) >= 3]

        if not tokens:
            continue
        if tokens[0] in {"semillas", "semilla"} and len(tokens) > 1:
            kw_set.add(tokens[1])
        elif tokens[0] in {"aceite", "harina", "leche", "manteca", "queso", "vino", "aceitunas"} and len(tokens) > 1:
            kw_set.add(tokens[0])
            kw_set.add(tokens[1])
        else:
            kw_set.add(tokens[0])

    return sorted(list(kw_set))


def scrape_criollo_robusto(targets: list[str], headed: bool = False) -> None:
    from playwright.sync_api import sync_playwright

    user = os.getenv("CRIOLLO_USER", "30714565091")
    password = os.getenv("CRIOLLO_PASSWORD", "lafinur3368")

    search_keywords = extract_keywords_from_targets(targets)
    LOG.info("Scraping exhaustivo de El Criollo. Total términos: %d", len(search_keywords))

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not headed)
        page = browser.new_page(viewport={"width": 1440, "height": 900})

        page.goto("https://tienda.distribuidoraelcriollo.com/Seguridad/Login", wait_until="domcontentloaded", timeout=60000)
        page.locator("#Usuario").fill(user)
        page.locator("#Clave").fill(password)
        page.get_by_role("button", name=re.compile(r"^\s*Ingresar\s*$", re.I)).click()
        page.wait_for_load_state("networkidle", timeout=30000)

        items: dict[str, dict[str, Any]] = {}

        for kw in search_keywords:
            search_url = f"https://tienda.distribuidoraelcriollo.com/Shop?search={quote_plus(kw)}"
            page.goto(search_url, wait_until="domcontentloaded", timeout=60000)

            while True:
                candidate_rows = page.locator("table tr:has(td)")
                try:
                    candidate_rows.first.wait_for(state="attached", timeout=3000)
                except Exception:
                    break

                count = candidate_rows.count()
                for i in range(count):
                    raw = candidate_rows.nth(i).evaluate("""row => {
                        const cells = Array.from(row.querySelectorAll(':scope > td'));
                        if (cells.length < 4) return null;
                        const link = cells[1].querySelector('a[title]');
                        return {
                            title: link ? link.getAttribute('title') : '',
                            detail: cells[2].innerText || '',
                            price: cells[3].innerText || ''
                        };
                    }""")
                    if not raw:
                        continue
                    description = str(raw["title"] or "").split("|", 1)[-1].strip()
                    price_text = str(raw["price"] or "").strip()
                    nominal = parse_price(price_text)
                    if not nominal or nominal < 50 or len(description) < 4:
                        continue

                    factor, base_unit, presentation = parse_presentation_by_unit(
                        f"{description} {str(raw['detail'])}", price_nominal=nominal
                    )

                    if "peso aprox" in str(raw["detail"]).lower() and not is_unit_only_item(normalize_text(description)):
                        factor, base_unit = 1.0, "KG"

                    key = normalize_text(description)
                    items[key] = {
                        "descripcion": description,
                        "precio_nominal": nominal,
                        "presentacion": presentation,
                        "factor_presentacion": factor,
                        "unidad_base": base_unit,
                        "precio_unitario_base": nominal / factor,
                    }

                next_btn = page.locator("a[rel='next'], li.next a, a:has-text('>')").first
                try:
                    if next_btn.is_visible(timeout=800):
                        next_btn.click()
                        page.wait_for_load_state("domcontentloaded", timeout=10000)
                    else:
                        break
                except Exception:
                    break

        browser.close()

    df = pd.DataFrame(list(items.values()))
    df.to_csv(CRIOLLO_CSV, sep=";", index=False, encoding="utf-8-sig")
    LOG.info("Catálogo El Criollo actualizado: %d productos", len(df))


def read_target_products(path: Path) -> list[str]:
    with open(path, "r", encoding="utf-8") as f:
        raw_lines = [line.strip() for line in f if line.strip() and not line.startswith("#")]

    cleaned, seen = [], set()
    for item in raw_lines:
        clean_key = re.sub(r"[^a-z0-9]+", " ", item.lower()).strip()
        canonical_name = CANONICAL_TARGET_MAP.get(clean_key, item)
        norm_key = normalize_text(canonical_name)
        if norm_key not in seen:
            seen.add(norm_key)
            cleaned.append(canonical_name)
    return cleaned


def read_posta(path: Path) -> pd.DataFrame:
    if not path.exists():
        LOG.warning(f"No se encontró {path.name}")
        return pd.DataFrame()
    frame = pd.read_excel(path, header=2, engine="openpyxl")
    prod_col = next(c for c in frame.columns if any(k in str(c).lower() for k in ["nombre", "descrip"]))
    price_col = next(c for c in frame.columns if "precio actual" in str(c).lower())
    um_col = next((c for c in frame.columns if str(c).strip().lower() in {"u.m.", "um", "unidad"}), None)

    df = pd.DataFrame()
    df["producto"] = frame[prod_col].fillna("").astype(str).str.strip()
    df["precio_nominal"] = frame[price_col].map(parse_price)
    df["unidad_lista"] = frame[um_col].fillna("UN").astype(str).str.strip().str.upper() if um_col else "UN"

    parsed = [parse_presentation_by_unit(p, u, pr) for p, u, pr in zip(df["producto"], df["unidad_lista"], df["precio_nominal"])]
    df[["factor_presentacion", "unidad_base", "presentacion"]] = pd.DataFrame(parsed, index=df.index)

    df["precio_unitario_base"] = df["precio_nominal"] / df["factor_presentacion"]
    df = df[df["producto"].ne("") & df["precio_unitario_base"].notna() & df["precio_unitario_base"].gt(0)].copy()
    df["_norm"] = df["producto"].map(normalize_text)
    return df.drop_duplicates("_norm").reset_index(drop=True)


def read_emporio_excel(path: Path) -> pd.DataFrame:
    if not path.exists():
        LOG.warning(f"No se encontró {path.name}")
        return pd.DataFrame()
    engine = "xlrd" if path.suffix.lower() == ".xls" else "openpyxl"
    raw = pd.read_excel(path, sheet_name=0, header=None, engine=engine)
    items = []

    # 1. Leer bloque superior (filas 0 a 45)
    for idx in range(min(45, len(raw))):
        prod_name = str(raw.iloc[idx, 0] or "").strip()
        extra_desc = str(raw.iloc[idx, 1] or "").strip()
        price_val = parse_price(raw.iloc[idx, 2])
        if not price_val or len(prod_name) < 3 or "ALMACEN" in prod_name or "PRECIOS" in prod_name:
            continue
        full_desc = f"{prod_name} {extra_desc}".strip()
        factor, base_unit, _ = parse_presentation_by_unit(full_desc, "UNI", price_val)
        items.append({
            "producto": full_desc,
            "precio_nominal": price_val,
            "factor_presentacion": factor,
            "unidad_base": base_unit,
            "precio_unitario_base": round(price_val / factor, 2),
            "_norm": normalize_text(full_desc),
        })

    # 2. Leer tabla principal
    header_idx = None
    for idx in range(25, min(65, len(raw))):
        row_str = " ".join([str(x) for x in raw.iloc[idx].values])
        if "ALMACEN" in row_str and "DESCRIPCION" in row_str:
            header_idx = idx
            break

    start_idx = (header_idx + 1) if header_idx is not None else 0
    table_df = raw.iloc[start_idx:].copy()

    for _, row in table_df.iterrows():
        prod_name = str(row[0] or "").strip()
        desc = str(row[1] or "").strip() if len(row) > 1 else ""
        unidad_str = str(row[2] or "").strip().upper() if len(row) > 2 else "UNI"
        price_val = parse_price(row[4] if len(row) > 4 else None) or parse_price(row[3] if len(row) > 3 else None) or parse_price(row[2] if len(row) > 2 else None)
        
        if not price_val or len(prod_name) < 3 or "PRECIOS VIGENTES" in prod_name or "ALMACEN" in prod_name:
            continue

        full_desc = f"{prod_name} {desc}".strip()
        factor, base_unit, _ = parse_presentation_by_unit(full_desc, unidad_str, price_val)
        items.append({
            "producto": full_desc,
            "precio_nominal": price_val,
            "factor_presentacion": factor,
            "unidad_base": base_unit,
            "precio_unitario_base": round(price_val / factor, 2),
            "_norm": normalize_text(full_desc),
        })

    df = pd.DataFrame(items)
    if not df.empty:
        df = df[df["precio_unitario_base"].gt(0)].drop_duplicates("_norm").reset_index(drop=True)
    return df


def read_criollo_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    frame = pd.read_csv(path, sep=";", encoding="utf-8-sig")
    df = pd.DataFrame()
    df["producto"] = frame["descripcion"].fillna("").astype(str).str.strip()
    df["precio_nominal"] = frame["precio_nominal"].astype(float)
    df["factor_presentacion"] = frame["factor_presentacion"].astype(float)
    df["unidad_base"] = frame["unidad_base"].astype(str)
    df["precio_unitario_base"] = frame["precio_unitario_base"].astype(float)
    df["_norm"] = df["producto"].map(normalize_text)
    return df.drop_duplicates("_norm").reset_index(drop=True)


def run_pricing_comparison(targets: list[str], df_posta: pd.DataFrame, df_criollo: pd.DataFrame, df_emporio: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    choices_p = df_posta["_norm"].tolist() if not df_posta.empty else []
    choices_c = df_criollo["_norm"].tolist() if not df_criollo.empty else []
    choices_e = df_emporio["_norm"].tolist() if not df_emporio.empty else []

    rows = []

    for item in sorted(targets, key=lambda s: s.lower()):
        norm = normalize_text(item)

        cands_p = find_all_safe_matches(norm, choices_p)
        cands_c = find_all_safe_matches(norm, choices_c)
        cands_e = find_all_safe_matches(norm, choices_e)

        rows_p = df_posta[df_posta["_norm"].isin(cands_p)].sort_values("precio_unitario_base").drop_duplicates("producto") if cands_p else pd.DataFrame()
        rows_c = df_criollo[df_criollo["_norm"].isin(cands_c)].sort_values("precio_unitario_base").drop_duplicates("producto") if cands_c else pd.DataFrame()
        rows_e = df_emporio[df_emporio["_norm"].isin(cands_e)].sort_values("precio_unitario_base").drop_duplicates("producto") if cands_e else pd.DataFrame()

        is_gen = is_generic_item(item)
        max_options = max(len(rows_p), len(rows_c), len(rows_e), 0)
        num_options = min(3, max_options) if is_gen else (1 if max_options > 0 else 0)

        if num_options == 0:
            rows.append({
                "Insumo": item,
                "Proveedor Ganador": "SIN MATCH CONFIABLE",
                "Producto Seleccionado": None,
                "Unidad Base": "UN",
                "Precio Base Ganador": None,
                "Ahorro Base ($)": 0.0,
                "Spread (%)": 0.0,
                "Precio Criollo": None,
                "Match Criollo": "-",
                "Precio Posta": None,
                "Match Posta": "-",
                "Precio Emporio": None,
                "Match Emporio": "-",
            })
            continue

        for opt_idx in range(num_options):
            matches = {}
            if opt_idx < len(rows_c):
                b = rows_c.iloc[opt_idx]
                matches["EL CRIOLLO"] = {"desc": b["producto"], "price": b["precio_unitario_base"], "unit": b["unidad_base"]}
            if opt_idx < len(rows_p):
                b = rows_p.iloc[opt_idx]
                matches["POSTA"] = {"desc": b["producto"], "price": b["precio_unitario_base"], "unit": b["unidad_base"]}
            if opt_idx < len(rows_e):
                b = rows_e.iloc[opt_idx]
                matches["EL NUEVO EMPORIO"] = {"desc": b["producto"], "price": b["precio_unitario_base"], "unit": b["unidad_base"]}

            if not matches:
                continue

            ganador = "SIN MATCH CONFIABLE"
            ahorro_val, spread = 0.0, 0.0
            p_ganador, p_alt = None, None
            prod_ganador, unidad_comun = None, "UN"

            units = {m["unit"] for m in matches.values()}
            if len(units) == 1:
                unidad_comun = list(units)[0]
                sorted_suppliers = sorted(matches.items(), key=lambda x: x[1]["price"])
                if len(sorted_suppliers) >= 2:
                    ganador = sorted_suppliers[0][0]
                    prod_ganador = sorted_suppliers[0][1]["desc"]
                    p_ganador = sorted_suppliers[0][1]["price"]
                    p_alt = sorted_suppliers[1][1]["price"]
                    ahorro_val = p_alt - p_ganador
                    spread = (ahorro_val / p_alt) * 100 if p_alt > 0 else 0.0
                else:
                    ganador = f"SOLO EN {sorted_suppliers[0][0]}"
                    prod_ganador = sorted_suppliers[0][1]["desc"]
                    p_ganador = sorted_suppliers[0][1]["price"]

            display_name = f"{item} (Opción {opt_idx + 1})" if num_options > 1 else item

            rows.append({
                "Insumo": display_name,
                "Proveedor Ganador": ganador,
                "Producto Seleccionado": prod_ganador,
                "Unidad Base": unidad_comun,
                "Precio Base Ganador": round(p_ganador, 2) if p_ganador else None,
                "Ahorro Base ($)": round(ahorro_val, 2),
                "Spread (%)": round(spread, 1),
                "Precio Criollo": round(matches["EL CRIOLLO"]["price"], 2) if "EL CRIOLLO" in matches else None,
                "Match Criollo": matches["EL CRIOLLO"]["desc"] if "EL CRIOLLO" in matches else "-",
                "Precio Posta": round(matches["POSTA"]["price"], 2) if "POSTA" in matches else None,
                "Match Posta": matches["POSTA"]["desc"] if "POSTA" in matches else "-",
                "Precio Emporio": round(matches["EL NUEVO EMPORIO"]["price"], 2) if "EL NUEVO EMPORIO" in matches else None,
                "Match Emporio": matches["EL NUEVO EMPORIO"]["desc"] if "EL NUEVO EMPORIO" in matches else "-",
            })

    df_matrix = pd.DataFrame(rows)
    df_recs = df_matrix[~df_matrix["Proveedor Ganador"].isin(["SIN MATCH CONFIABLE", "UNIDADES INCOMPARABLES"])].copy()
    return df_recs, df_matrix


def export_excel(df_matrix: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df_matrix.to_excel(writer, sheet_name="Comparativa_General", index=False)

    wb = load_workbook(path)
    ws = wb.active
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 28

    wb.save(path)
    LOG.info("Reporte Excel guardado: %s", path.name)


def send_email_report(user: str, password: str, recipient: str, excel_path: Path, df_matrix: pd.DataFrame) -> None:
    LOG.info("Enviando correo a %s...", recipient)
    rows_html = ""

    for _, r in df_matrix.iterrows():
        u = r["Unidad Base"] or "UN"
        p_c = f"${r['Precio Criollo']:,.2f}/{u}" if pd.notna(r["Precio Criollo"]) else "-"
        p_p = f"${r['Precio Posta']:,.2f}/{u}" if pd.notna(r["Precio Posta"]) else "-"
        p_e = f"${r['Precio Emporio']:,.2f}/{u}" if pd.notna(r["Precio Emporio"]) else "-"

        ahorro_str = f"${r['Ahorro Base ($)']:,.2f}/{u} ({r['Spread (%)']}%)" if r["Ahorro Base ($)"] > 0 else "-"

        desc_c = html.escape(str(r["Match Criollo"]))
        desc_p = html.escape(str(r["Match Posta"]))
        desc_e = html.escape(str(r["Match Emporio"]))

        rows_html += f"""
        <tr style="border-bottom: 1px solid #e0e0e0;">
            <td style="padding: 10px 8px; font-weight: bold; color: #1a237e;">{html.escape(str(r['Insumo']))}</td>
            <td style="padding: 10px 8px; font-weight: bold; color: #00695c;">{html.escape(str(r['Proveedor Ganador']))}</td>
            <td style="padding: 10px 8px; text-align: right; background-color: #fafafa;">
                <b>{p_c}</b><br><span style="font-size: 10px; color: #616161;">{desc_c}</span>
            </td>
            <td style="padding: 10px 8px; text-align: right;">
                <b>{p_p}</b><br><span style="font-size: 10px; color: #616161;">{desc_p}</span>
            </td>
            <td style="padding: 10px 8px; text-align: right; background-color: #fafafa;">
                <b>{p_e}</b><br><span style="font-size: 10px; color: #616161;">{desc_e}</span>
            </td>
            <td style="padding: 10px 8px; text-align: right; color: #2e7d32; font-weight: bold;">{ahorro_str}</td>
        </tr>
        """

    html_content = f"""
    <html>
    <body style="font-family: Arial, sans-serif; color: #333; margin: 20px;">
        <h2 style="color: #1F4E78; border-bottom: 2px solid #1F4E78; padding-bottom: 8px;">
            Comparativa de Precios Multi-Proveedor: El Criollo | Posta | El Nuevo Emporio
        </h2>
        <p style="font-size: 13px; color: #555;">
            Orden alfabético. Marcas y variedades normalizadas. Fideos, puré y tapas cotizados por unidad de paquete/lata.
        </p>
        <table style="border-collapse: collapse; width: 100%; max-width: 1050px; font-size: 12px;">
            <thead>
                <tr style="background-color: #1F4E78; color: white;">
                    <th style="padding: 10px 8px; text-align: left;">Insumo Requerido</th>
                    <th style="padding: 10px 8px; text-align: left;">Mejor Opción</th>
                    <th style="padding: 10px 8px; text-align: right;">El Criollo</th>
                    <th style="padding: 10px 8px; text-align: right;">Posta</th>
                    <th style="padding: 10px 8px; text-align: right;">El Nuevo Emporio</th>
                    <th style="padding: 10px 8px; text-align: right;">Ahorro Unitario</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>
    </body>
    </html>
    """

    msg = email.message.EmailMessage()
    msg["Subject"] = f"Reporte Pricing Tri-Proveedor - {datetime.now().strftime('%d/%m/%Y')}"
    msg["From"] = user
    msg["To"] = recipient
    msg.set_content("Adjunto encontrarás el reporte detallado.")
    msg.add_alternative(html_content, subtype="html")

    with open(excel_path, "rb") as f:
        msg.add_attachment(f.read(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=excel_path.name)

    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
        server.login(user, password)
        server.send_message(msg)

    LOG.info("Mail enviado con éxito.")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--no-mail", action="store_true")
    parser.add_argument("--skip-scrape", action="store_true")
    args = parser.parse_args()

    targets = read_target_products(LISTA_BUSQUEDA)
    if not targets:
        LOG.error("productos a buscar.txt está vacío.")
        return 1

    if not args.skip_scrape:
        if CRIOLLO_CSV.exists():
            CRIOLLO_CSV.unlink()
        scrape_criollo_robusto(targets, headed=args.headed)

    df_posta = read_posta(POSTA_XLSX)
    df_emporio = read_emporio_excel(EMPORIO_XLS)
    df_criollo = read_criollo_csv(CRIOLLO_CSV)

    _, df_matrix = run_pricing_comparison(targets, df_posta, df_criollo, df_emporio)
    export_excel(df_matrix, OUTPUT_XLSX)

    if not args.no_mail:
        user = os.getenv("GMAIL_USER", "aaron.derek0@gmail.com")
        pw = os.getenv("GMAIL_APP_PASSWORD", "qbkt iqzu wpwh gquj")
        dest = os.getenv("DESTINATARIO_MAIL", user)
        send_email_report(user, pw.replace(" ", ""), dest, OUTPUT_XLSX, df_matrix)

    return 0


if __name__ == "__main__":
    sys.exit(main())