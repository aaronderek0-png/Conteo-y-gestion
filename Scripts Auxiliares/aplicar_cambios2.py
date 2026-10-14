import os
from pathlib import Path
os.chdir(r"C:\Users\EXO G12\Desktop\Agente Mishiguene")


def edit(path, pairs):
    s = open(path, encoding="utf-8").read()
    for a, b in pairs:
        assert a in s, (path, a[:70])
        s = s.replace(a, b, 1)
    open(path, "w", encoding="utf-8").write(s)


# ---- Carpeta Revision: se arma escaneando las copias (fuente única, idempotente)
s = open("procesar_lote_wpp.py", encoding="utf-8").read()
i = s.index("def escribir_carpeta_revision")
j = s.index("def main(argv")
nuevo = '''def escribir_carpeta_revision(filas: list[dict], analisis: dict, cl: dict) -> Path:
    """'Carpeta Revision/Revision Facturas IA.xlsx': espejo de TODAS las filas pintadas por el agente que siguen en las
    copias (amarillo #FFF8E1 = OK pre-aprobada, rojo #FFEBEE = a revisar) + alertas por comprobante."""
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    carpeta = BASE_DIR / "Carpeta Revision"
    carpeta.mkdir(exist_ok=True)
    nuevas = {(f["local"], f["fila_excel"]): f for f in filas}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Filas inyectadas"
    ws.append(["Local", "Fila en la copia", "N° factura", "Fecha", "Proveedor", "Presentación / marca", "Cantidad", "Precio",
               "Desc.", "Estado", "Nota o motivo", "Precio transferencia P_UN"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for copia in sorted((imp.COPIAS_DIR).glob("*_FACTURAS_COPIA.xlsx")):
        local = copia.name.replace("_FACTURAS_COPIA.xlsx", "")
        hoja = openpyxl.load_workbook(copia)["Facturas"]
        for r in range(2, hoja.max_row + 1):
            col = (hoja.cell(r, 1).fill.fgColor.rgb or "") if hoja.cell(r, 1).fill and hoja.cell(r, 1).fill.fill_type else ""
            col = str(col)[-6:].upper()
            if col not in ("FFF8E1", "FFEBEE") or hoja.cell(r, 1).value in (None, ""):
                continue
            n = nuevas.get((local, r)) or {}
            rojo = col == "FFEBEE"
            ws.append([local, r, hoja.cell(r, 1).value, hoja.cell(r, 2).value, hoja.cell(r, 4).value, hoja.cell(r, 5).value,
                       hoja.cell(r, 7).value, hoja.cell(r, 8).value, hoja.cell(r, 9).value,
                       "REVISAR" if rojo else "OK (pre-aprobada)",
                       hoja.cell(r, 6).value or "; ".join((n.get("motivos") or "",)), n.get("transferencia")])
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=col)
    wa = wb.create_sheet("Alertas por comprobante")
    wa.append(["Archivo", "Tipo", "N°", "Proveedor", "Local", "Regla de local", "Alerta"])
    for c in cl["comprobantes"]:
        an = analisis.get(c["archivo"])
        if not an:
            continue
        for al in an["alertas"]:
            if al:
                wa.append([c["archivo"], c.get("tipo"), c.get("numero_factura"), c.get("proveedor_detectado"),
                           an["local"], an["por"], al])
    destino = carpeta / "Revision Facturas IA.xlsx"
    wb.save(destino)
    return destino


'''
s = s[:i] + nuevo + s[j:]
s = s.replace('''    if escribir and escrituras:
        escribir_carpeta_revision(revision_filas, analisis, cl)''', '''    if escribir:
        escribir_carpeta_revision(revision_filas, analisis, cl)''', 1)
open("procesar_lote_wpp.py", "w", encoding="utf-8").write(s)

# ---- datos_operativos: carpeta por defecto
s = open("datos_operativos.py", encoding="utf-8").read()
s = s.replace('default Facturas_WPP/', "default 'Facturas Mishiguene/'")
s = s.replace('origen: Path = BASE_DIR / "Facturas_WPP"', 'origen: Path = BASE_DIR / "Facturas Mishiguene"')
s = s.replace('else BASE_DIR / "Facturas_WPP"', 'else BASE_DIR / "Facturas Mishiguene"')
open("datos_operativos.py", "w", encoding="utf-8").write(s)
print("ok")
