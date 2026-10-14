import shutil
from pathlib import Path
import os
os.chdir(r"C:\Users\EXO G12\Desktop\Agente Mishiguene")


def edit(path, pairs):
    s = open(path, encoding="utf-8").read()
    for a, b in pairs:
        assert a in s, (path, a[:70])
        s = s.replace(a, b, 1)
    open(path, "w", encoding="utf-8").write(s)


# ---- analisis_gestion_comercial: Remitos Chimidoha + matriz de recargos
edit("analisis_gestion_comercial.py", [(
'''    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s", datefmt="%H:%M:%S")
    resultados: list[dict[str, Any]] = []''',
'''    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s", datefmt="%H:%M:%S")
    # Pasamano Obrador -> Chimidoha/locales: matriz de recargos y aprendizaje de la hoja 'Remitos Chimidoha' (si ya existe)
    import transferencias_obrador as tr
    remitos_chimi = tr.buscar_remitos_chimidoha()
    if remitos_chimi:
        apr = tr.aprender_de_remitos_chimidoha(remitos_chimi)
        LOG.info("Remitos Chimidoha: %s -> %s", [p.name for p in remitos_chimi], apr.get("por_tipo") or apr.get("estado"))
    else:
        LOG.info("Remitos Chimidoha: no hay archivo todavia; transferencias con la matriz por defecto "
                 "(panificados/descartables x1.20, carnes/pescados x1.10, almacen al costo, elaborados por BOM)")
    resultados: list[dict[str, Any]] = []''')])

# ---- imputacion_facturas
edit("imputacion_facturas.py", [
('''    rojo = PatternFill("solid", fgColor="FFC7CE")''',
 '''    rojo = PatternFill("solid", fgColor="FFEBEE")         # requiere revisión (rojo suave)
    ok_fill = PatternFill("solid", fgColor="FFF8E1")      # inyectada OK, pendiente de aprobación (amarillo suave)'''),
('''            for c, v in valores.items():
                ws.cell(r, c).value = v
                if reviso:
                    ws.cell(r, c).fill = rojo''',
 '''            for c, v in valores.items():
                ws.cell(r, c).value = v
            for c in range(1, 10):                      # toda fila nueva va pintada hasta que alguien la apruebe
                ws.cell(r, c).fill = rojo if reviso else ok_fill'''),
('''BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "Arcos_data"''',
 '''BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "Arcos_data"
CARPETA_ENTRADA = BASE_DIR / "Facturas Mishiguene"          # carpeta diaria de comprobantes (sin guiones)
CARPETA_ENTRADA_VIEJA = BASE_DIR / "Facturas_WPP"           # lectura retrocompatible
CARPETA_REVISION = BASE_DIR / "Carpeta Revision"            # espejo de lo inyectado (OK / a revisar)
COLOR_OK = "FFF8E1"                                         # amarillo suave: inyectada, falta aprobar
COLOR_REVISAR = "FFEBEE"                                    # rojo suave: requiere revisión humana'''),
('''    ap_.add_argument("entrada", nargs="?", help="comprobante .json/.pdf/.jpg/... o carpeta")''',
 '''    ap_.add_argument("entrada", nargs="?", help="comprobante .json/.pdf/.jpg/... o carpeta "
                                                "(default: 'Facturas Mishiguene/')")'''),
('''    else:
        ap_.print_help()
        return 1
''', '''    else:
        carpeta = CARPETA_ENTRADA if CARPETA_ENTRADA.is_dir() else CARPETA_ENTRADA_VIEJA
        if not carpeta.is_dir():
            ap_.print_help()
            return 1
        return main([str(carpeta)] + [a for a in (argv if argv is not None else sys.argv[1:])])
'''),
('''    ws.update(f"A{primera}:I{primera + len(fmt) - 1}", fmt, value_input_option="USER_ENTERED")
    return {"escrito": True, "desde_fila": primera, "filas": len(fmt)}''',
 '''    ws.update(f"A{primera}:I{primera + len(fmt) - 1}", fmt, value_input_option="USER_ENTERED")
    # amarillo suave hasta que alguien las apruebe con el menú de GAS/aprobador_facturas_ia.gs
    ws.format(f"A{primera}:I{primera + len(fmt) - 1}",
              {"backgroundColor": {"red": 1.0, "green": 0.973, "blue": 0.882}})
    return {"escrito": True, "desde_fila": primera, "filas": len(fmt)}'''),
('''    return {"proveedor_oficial": proveedor, "modo": "PRORRATEO" if items is not comp.get("items") else "FILAS",
            "filas": filas, "notas": notas}''',
 '''    _anotar_transferencia(filas, loc)
    return {"proveedor_oficial": proveedor, "modo": "PRORRATEO" if items is not comp.get("items") else "FILAS",
            "filas": filas, "notas": notas}'''),
('''# ---- publicación en la planilla real (solo con confirmación) -----------------''',
 '''def _anotar_transferencia(filas: list[dict], loc: "Local") -> None:
    """Congreso (Obrador) compra y luego 'vende' a Chimidoha y a los locales: cada fila INYECTAR lleva el precio de
    transferencia por la matriz (panificados/descartables x1.20, carnes/pescados x1.10, almacén al costo)."""
    if loc.nombre != "Congreso":
        return
    import transferencias_obrador as tr
    for f in filas:
        if f.get("accion") != "INYECTAR" or not f.get("presentacion") or not f.get("P_UN"):
            continue
        cat = loc.catalogo[loc.catalogo["PRESENTACION"] == f["presentacion"]]
        if cat.empty:
            continue
        cat2 = cat.iloc[0]["CAT2"]
        f["transferencia"] = {"tipo": tr.tipo_transferencia(cat2), "factor": tr.factor_transferencia(cat2),
                              "p_un_transferencia": tr.precio_transferencia(f["P_UN"], cat2)}


# ---- publicación en la planilla real (solo con confirmación) -----------------'''),
])

# ---- invoice_parser
edit("invoice_parser.py", [(
'''        alertas.append("RAZÓN SOCIAL CHIMI DOHA con sello/nota Cerviño: verificar si es gasto del café (Chimidoha)")''',
'''        alertas.append("INFO compra centralizada: razón social Chimi Doha pero sello Cerviño → Congreso (Obrador) compra "
                       "y luego transfiere a Chimidoha con la matriz de recargos")''')])

# ---- procesar_lote_wpp
edit("procesar_lote_wpp.py", [
('''    ap_.add_argument("carpeta", nargs="?", default=str(BASE_DIR / "Facturas_WPP"))''',
 '''    ap_.add_argument("carpeta", nargs="?", default=str(carpeta_por_defecto()))'''),
('''def sha(path: Path) -> str:''',
 '''def carpeta_por_defecto() -> Path:
    """'Facturas Mishiguene/' (sin guiones); si todavía no existe se lee la carpeta vieja 'Facturas_WPP/'."""
    nueva, vieja = BASE_DIR / "Facturas Mishiguene", BASE_DIR / "Facturas_WPP"
    return nueva if nueva.is_dir() or not vieja.is_dir() else vieja


def sha(path: Path) -> str:'''),
('''    if m0 and regla["local"] and m0["local"] != regla["local"]:''',
 '''    if local is None and "Chimidoha" in locales and "chimi" in imp.norm(comp.get("razon_social_cliente")):
        if imp.norm(comp.get("proveedor_detectado")) in locales["Chimidoha"]._idx_prov or \\
                imp.norm(comp.get("proveedor_planilla")) in locales["Chimidoha"]._idx_prov:
            local, por = "Chimidoha", "proveedor con historial cargado directo en Chimidoha y razón social Chimi"
    if m0 and regla["local"] and m0["local"] != regla["local"]:'''),
('''        alertas.append("MULTILOCAL: " + extra["multilocal"] + " → prorratear entre los locales de cada día")''',
 '''        alertas.append("MULTILOCAL (100 % en Congreso, etiqueta MULTILOCAL): " + extra["multilocal"])'''),
('''            escrituras[nombre] = {"copia": loc.copia.name, "filas": rep["escritas"], "verificadas": rep["verificadas"],
                                  "marcadas": len(marcas_por_local.get(nombre, []))}''',
 '''            escrituras[nombre] = {"copia": loc.copia.name, "filas": rep["escritas"], "verificadas": rep["verificadas"],
                                  "marcadas": len(marcas_por_local.get(nombre, []))}
            for r_, f_ in rep["filas"]:
                revision_filas.append({"local": nombre, "fila_excel": r_, "N": f_.get("A"), "fecha": f_.get("B"),
                                       "proveedor": f_.get("D"), "presentacion": f_.get("presentacion"),
                                       "original": f_.get("descripcion_original"), "G": f_.get("G") or f_.get("cant_orig"),
                                       "H": f_.get("H") or f_.get("precio_orig"), "I": f_.get("I"),
                                       "estado": "REVISAR" if f_.get("accion") == "REVISAR" else "OK (pre-aprobada)",
                                       "motivos": "; ".join((f_.get("motivos") or []) + (f_.get("avisos") or [])),
                                       "transferencia": (f_.get("transferencia") or {}).get("p_un_transferencia")})'''),
('''    escrituras = {}
    if escribir:''',
 '''    escrituras = {}
    revision_filas: list[dict] = []
    if escribir:'''),
('''    pushes = {}
    if args.push_drive:''',
 '''    if escribir and escrituras:
        escribir_carpeta_revision(revision_filas, analisis, cl)
    pushes = {}
    if args.push_drive:'''),
('''def main(argv: list[str] | None = None) -> int:''',
 '''def escribir_carpeta_revision(filas: list[dict], analisis: dict, cl: dict) -> Path:
    """Refleja en 'Carpeta Revision/' lo inyectado (amarillo #FFF8E1 OK / rojo #FFEBEE a revisar) y las alertas."""
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    carpeta = BASE_DIR / "Carpeta Revision"
    carpeta.mkdir(exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Filas inyectadas"
    ws.append(["Local", "Fila en la copia", "N° factura", "Fecha", "Proveedor", "Presentación", "Texto original",
               "Cantidad", "Precio", "Desc.", "Estado", "Motivos / avisos", "Precio transferencia P_UN"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append([f["local"], f["fila_excel"], f["N"], f["fecha"], f["proveedor"], f["presentacion"], f["original"],
                   f["G"], f["H"], f["I"], f["estado"], f["motivos"], f["transferencia"]])
        color = "FFEBEE" if f["estado"] == "REVISAR" else "FFF8E1"
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=color)
    wa = wb.create_sheet("Alertas por comprobante")
    wa.append(["Archivo", "Tipo", "N°", "Proveedor", "Local", "Regla de local", "Alerta"])
    for c in cl["comprobantes"]:
        an = analisis.get(c["archivo"])
        if not an:
            continue
        for al in an["alertas"]:
            if al:
                wa.append([c["archivo"], c.get("tipo"), c.get("numero_factura"), c.get("proveedor_detectado"),
                           an["local"], an["por"], al])
    destino = carpeta / "Revision Facturas IA.xlsx"
    wb.save(destino)
    return destino


def main(argv: list[str] | None = None) -> int:'''),
])

# ---- migración de carpeta (el cache va por SHA-256: no se duplica nada)
src, dst = Path("Facturas_WPP"), Path("Facturas Mishiguene")
dst.mkdir(exist_ok=True)
Path("Carpeta Revision").mkdir(exist_ok=True)
movidos = 0
if src.is_dir():
    for p in src.iterdir():
        if p.is_file() and not (dst / p.name).exists():
            shutil.move(str(p), str(dst / p.name))
            movidos += 1
print("ediciones ok; archivos migrados:", movidos)
