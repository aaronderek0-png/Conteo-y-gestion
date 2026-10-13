"""Motor de Recetas (BOM) y puente de códigos para el ecosistema Mishiguene.

Fuente: 'Recetas' (Drive: "Recetas 99% gas", exportada a Copias_Gestion/originales_export/Recetas.xlsx)
    Insumos   : A=SKU  B=Producto  C=Tipo  D=Rubro  E=UNIDAD_M (g/ml/un/kg)  F=P_UN (precio por esa unidad)
    Recetas   : A=COD_RECETA  B=NOM_RECETA  C=NOMBRE_RECETA  D=COD (insumo o sub-receta)  E=NOM_INSUMO
                F=CANTIDAD  G=UNIDAD  H=P_UN  K=REND (rendimiento del batch)  L=C_UN  M=Activo
    Costos    : A=COD_REC  B=NOM_RECETA  C=NOMBRE  D=C_UN  H=U. MEDIDA (unidad de la receta: gr/un/porc)
    Obrador   : A=SKU  C=COSTO  H=MARGEN ELEGIDO  I=P_TRANSFERENCIA
    Equivalencias: I=CODIGO M (Maxirest)  J=NOMBRE M  K=OBRADOR  L=NOMBRE RECETA  M=CODIGO RECETA

Reglas:
  * Elaboración = código que empieza con M / MR / MO / OBR / P0, o nombre con 'MADRE'. Se explota como receta
    (nunca como insumo simple).
  * factor de escala = cantidad requerida / REND; cada línea se multiplica por ese factor.
  * g, gr, ml, cc -> base 'g/ml' (se valoriza por KG/LT al multiplicar el precio por 1000); kg/lt/l en la línea
    se pasan a g/ml (x1000); un/porc/caja -> 1:1.
  * Un insumo cuyo UNIDAD_M es kg/lt se pasa a precio por g/ml (/1000).
"""
from __future__ import annotations

import csv
import difflib
import hashlib
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from defensivo import (SchemaError, detectar_ciclos, es_panificado_comprado, finito, parse_numero,
                       resolver_esquema, safe_div)

MAX_PROFUNDIDAD = 40   # niveles de sub-recetas anidadas tolerados antes de cortar
ELAB_RE = re.compile(r"^(M\d|MR\d|MO\d|OBR\d|P0\d)|\bMADRE\b", re.I)
UNIDADES_MASA = {"g", "gr", "grs", "gramo", "gramos", "ml", "cc", "cm3"}
UNIDADES_X1000 = {"kg", "kgs", "kilo", "kilos", "lt", "lts", "l", "litro", "litros"}
UNIDADES_ENTERAS = {"un", "u", "unidad", "unidades", "porc", "porcion", "porción", "caja", "bolsa", "pack"}


def es_elaboracion(codigo: str, nombre: str = "") -> bool:
    return bool(ELAB_RE.search(str(codigo or "").strip()) or re.search(r"\bMADRE\b", str(nombre or ""), re.I))


def _n(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _base(unidad) -> tuple[str, float]:
    """-> (clase 'masa'|'un', multiplicador para llevar la cantidad de la línea a g/ml o un)."""
    u = str(unidad or "").strip().lower().rstrip(".")
    if u in UNIDADES_X1000:
        return "masa", 1000.0
    if u in UNIDADES_MASA:
        return "masa", 1.0
    return "un", 1.0


@dataclass
class Resultado:
    codigo: str
    cantidad: float
    unidad: str                       # unidad de la receta: 'gr' | 'un' | 'porc'
    costo_total: float                # costo de producir `cantidad`
    costo_unitario: float             # por gr/un (antes de x1000)
    precio_gestion: float             # por KG/LT si gr/ml; 1:1 si un/porc
    primarios: dict = field(default_factory=dict)     # SKU insumo -> cantidad requerida (g/ml/un)
    faltantes: list = field(default_factory=list)     # insumos sin precio / sin definición
    factor: float = 1.0

    @property
    def completo(self) -> bool:
        return not self.faltantes


class MotorRecetas:
    def __init__(self, xlsx: Path):
        """Lee el libro por NOMBRE de encabezado (no por índice). Falta un encabezado obligatorio -> SchemaError."""
        wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
        self.avisos: list[str] = []

        def hoja(nombre, spec, obligatorios):
            ws = wb[nombre] if nombre in wb.sheetnames else None
            if ws is None:
                raise SchemaError(f"Recetas.xlsx: falta la hoja '{nombre}'")
            cab = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
            cols, hdr, av = resolver_esquema(cab, spec, obligatorios=obligatorios, nombre=f"Recetas/{nombre}")
            self.avisos += av
            return ws, cols, hdr

        def g(r, cols, k, default=None):
            i = cols.get(k)
            return r[i] if i is not None and i < len(r) else default

        # --- Insumos
        ws, ci, hi = hoja("Insumos", {"sku": ["SKU"], "nombre": ["PRODUCTO"], "tipo": ["TIPO"], "rubro": ["RUBRO / CATEGORIA"],
                                      "unidad": ["UNIDAD_M"], "precio": ["P_UN"]}, ("sku", "unidad", "precio"))
        self.insumos: dict[str, dict] = {}
        for r in ws.iter_rows(min_row=hi + 1, values_only=True):
            sku = g(r, ci, "sku")
            if not sku:
                continue
            u = str(g(r, ci, "unidad") or "").strip().lower()
            p_, _ = parse_numero(g(r, ci, "precio"))
            if p_ is not None and p_ < 0:
                self.avisos.append(f"Insumo {sku}: precio negativo {p_} descartado")
                p_ = None
            if p_ is not None and u in UNIDADES_X1000:
                p_, u = p_ / 1000.0, "g" if u in ("kg", "kgs", "kilo", "kilos") else "ml"
            self.insumos[str(sku).strip().upper()] = dict(sku=str(sku).strip(), nombre=g(r, ci, "nombre"),
                                                          tipo=g(r, ci, "tipo"), rubro=g(r, ci, "rubro"),
                                                          unidad=u, precio=p_)
        # --- Recetas
        ws, cr, hr = hoja("Recetas", {"cod": ["COD_RECETA"], "nom": ["NOM_RECETA"], "nombre": ["NOMBRE_RECETA"],
                                      "ins": ["COD"], "insnom": ["NOM_INSUMO"], "cant": ["CANTIDAD"], "uni": ["UNIDAD"],
                                      "rend": ["REND"], "cun": ["C_UN"]}, ("cod", "ins", "cant", "uni", "rend", "cun"))
        self.col_recetas_cun = cr["cun"]
        self.recetas: dict[str, dict] = {}
        self.por_nom: dict[str, str] = {}
        self.filas_recetas: dict[str, list[int]] = {}      # COD_RECETA -> filas de la hoja Recetas
        for n_fila, r in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), start=hr + 1):
            cod_ = g(r, cr, "cod")
            if not cod_:
                continue
            cod = str(cod_).strip().upper()
            self.filas_recetas.setdefault(cod, []).append(n_fila)
            rend, _ = parse_numero(g(r, cr, "rend"))
            if rend is not None and rend <= 0:
                rend = None
            cun, _ = parse_numero(g(r, cr, "cun"))
            rc = self.recetas.setdefault(cod, dict(cod=cod, nom=str(g(r, cr, "nom") or "").strip().upper(),
                                                   nombre=g(r, cr, "nombre"), rend=rend, c_un_hoja=cun, lineas=[]))
            if rc["rend"] is None:
                rc["rend"] = rend
            ins = g(r, cr, "ins")
            if ins:
                cant, av = parse_numero(g(r, cr, "cant"))
                if cant is None or cant < 0:
                    self.avisos.append(f"Receta {cod}: cantidad no válida {g(r, cr, 'cant')!r} para {ins} "
                                       f"(la receta queda incompleta)")
                    cant = None
                elif av.startswith("texto con unidad"):
                    self.avisos.append(f"Receta {cod}: cantidad '{g(r, cr, 'cant')}' de {ins} traía unidad en el texto")
                rc["lineas"].append((str(ins).strip().upper(), g(r, cr, "insnom"), cant, g(r, cr, "uni")))
            if rc["nom"]:
                self.por_nom.setdefault(rc["nom"], cod)
        # --- Costos (unidad de cada receta)
        ws, cc, hc = hoja("Costos", {"cod": ["COD_REC"], "cun": ["C_UN"], "uni": ["U. MEDIDA"]}, ("cod", "cun", "uni"))
        self.col_costos_cun = cc["cun"]
        self.unidad_receta: dict[str, str] = {}
        self.fila_costos: dict[str, int] = {}               # COD_REC -> fila de la hoja Costos
        for n_fila, r in enumerate(ws.iter_rows(min_row=hc + 1, values_only=True), start=hc + 1):
            c_ = g(r, cc, "cod")
            if c_:
                k = str(c_).strip().upper()
                self.fila_costos[k] = n_fila
                self.unidad_receta[k] = str(g(r, cc, "uni") or "").strip().lower()
        self._idx_nombres = None
        # --- Obrador (lista de transferencia)
        self.obrador: dict[str, dict] = {}
        if "Obrador" in wb.sheetnames:
            try:
                ws, co, ho = hoja("Obrador", {"sku": ["SKU"], "costo": ["COSTO"], "margen": ["MARGEN ELEGIDO"],
                                              "transf": ["P_TRANSFERENCIA"]}, ("sku", "costo"))
                for r in ws.iter_rows(min_row=ho + 1, values_only=True):
                    if g(r, co, "sku"):
                        self.obrador[str(g(r, co, "sku")).strip().upper()] = dict(
                            costo=parse_numero(g(r, co, "costo"))[0], margen=parse_numero(g(r, co, "margen"))[0],
                            transferencia=parse_numero(g(r, co, "transf"))[0])
            except SchemaError as e:
                self.avisos.append(f"Obrador omitida: {e}")
        # --- Equivalencias Maxirest (numérico) -> receta
        self.equiv_maxirest: dict[str, str] = {}
        if "Equivalencias" in wb.sheetnames:
            try:
                ws, ce, he = hoja("Equivalencias", {"cm": ["CODIGO M"], "obr": ["OBRADOR"], "rec": ["CODIGO RECETA"]},
                                  ("cm",))
                for r in ws.iter_rows(min_row=he + 1, values_only=True):
                    cm = g(r, ce, "cm")
                    dest = g(r, ce, "rec") or g(r, ce, "obr")
                    if cm is not None and dest:
                        k = str(int(cm)) if isinstance(cm, (int, float)) else str(cm)
                        self.equiv_maxirest[k] = str(dest).strip().upper()
            except SchemaError as e:
                self.avisos.append(f"Equivalencias omitida: {e}")
        wb.close()
        # --- integridad del grafo BOM
        grafo = {c: [self.resolver_receta(l[0]) for l in rc["lineas"]
                     if self.resolver_receta(l[0]) and self._es_sub(l[0], l[1])] for c, rc in self.recetas.items()}
        self.ciclos = detectar_ciclos(grafo)
        self.en_ciclo = {n for c in self.ciclos for n in c}
        self.sin_rend = sorted(c for c, rc in self.recetas.items() if not rc["rend"])

    def _es_sub(self, sku: str, nom) -> bool:
        if es_panificado_comprado(sku):
            return False
        sub = self.resolver_receta(sku)
        if sub is None:
            return False
        ins = self.insumos.get(str(sku).upper())
        return not (ins is not None and ins["precio"] is not None and not es_elaboracion(sku, nom))

    # ------------------------------------------------------------------ puente de códigos
    def resolver_receta(self, codigo: str) -> str | None:
        """COD_RECETA directo; si no, NOM_RECETA (código corto tipo POBAKLA/BAMAPOLA) -> COD_RECETA."""
        c = str(codigo or "").strip().upper()
        if c in self.recetas:
            return c
        return self.por_nom.get(c)

    def canonico(self, codigo: str) -> str | None:
        """SKU canónico (NOM_RECETA) que usa el Maestro de Precios para una receta identificada por COD_RECETA."""
        c = self.resolver_receta(codigo)
        return self.recetas[c]["nom"] if c and self.recetas[c]["nom"] else None

    def _sub_receta(self, sku: str, nom) -> str | None:
        """Una línea es sub-receta si su código es una receta (COD_RECETA o NOM_RECETA); si además existe como
        insumo con precio, gana la receta solo cuando es elaboración (M/MR/MO/OBR/MADRE).
        Los panificados comprados (PITA, BAGEL...) son insumo final: no se explotan (regla del GAS)."""
        if es_panificado_comprado(sku):
            return None
        sub = self.resolver_receta(sku)
        if sub is None:
            return None
        ins = self.insumos.get(sku)
        if ins is not None and ins["precio"] is not None and not es_elaboracion(sku, nom):
            return None
        return sub

    # ------------------------------------------------------------------ BOM
    def _costo_unit(self, cod: str, pila: tuple, memo: dict):
        """-> (costo por gr/ml/un de la receta, faltantes). Recursivo con guardas de ciclo."""
        if cod in memo:
            return memo[cod]
        rc = self.recetas.get(cod)
        if rc is None:
            return None, [f"{cod}: receta inexistente"]
        if cod in pila or len(pila) > MAX_PROFUNDIDAD:
            return None, [f"{cod}: ciclo o profundidad > {MAX_PROFUNDIDAD} en la cadena {' > '.join(pila[-4:])}"]
        if not rc["rend"]:
            return None, [f"{cod}: sin REND (<=0 o vacío)"]
        total, falt = 0.0, []
        for sku, nom, cant, uni in rc["lineas"]:
            if cant is None:
                falt.append(f"{sku}: cantidad vacía o inválida")
                continue
            clase, mult = _base(uni)
            q = cant * mult
            sub = self._sub_receta(sku, nom)
            if sub and sub != cod:
                cu, f = self._costo_unit(sub, pila + (cod,), memo)
                falt += f
                if cu is None:
                    continue
                # la sub-receta se costea por su unidad (gr/un); la línea viene en esa misma unidad
                total += cant * cu
                continue
            ins = self.insumos.get(sku)
            if ins is None or ins["precio"] is None:   # precio 0 es válido (p. ej. AGUA)
                falt.append(f"{sku} ({nom})")
                continue
            total += q * ins["precio"]
        cu = safe_div(total, rc["rend"])
        res = (cu, falt if cu is not None else falt + [f"{cod}: costo no finito"])
        memo[cod] = res
        return res

    def explotar(self, codigo: str, cantidad: float = 1.0) -> Resultado | None:
        """Explota la receta `codigo` para `cantidad` (en la unidad de la receta). Factor = cantidad / REND."""
        cod = self.resolver_receta(codigo)
        if cod is None:
            return None
        rc = self.recetas[cod]
        cu, falt = self._costo_unit(cod, (), {})
        if cu is None:
            return Resultado(cod, cantidad, "", 0.0, 0.0, 0.0, faltantes=falt)
        unidad = self.unidad_receta.get(cod) or "gr"
        factor = cantidad / rc["rend"]
        prim: dict[str, float] = {}

        def acumular(c, f, pila):
            for sku, nom, cant, uni in self.recetas[c]["lineas"]:
                if cant is None:
                    continue
                _, mult = _base(uni)
                sub = self._sub_receta(sku, nom)
                if sub and sub != c and sub not in pila:
                    acumular(sub, f * cant / (self.recetas[sub]["rend"] or 1), pila + (c,))
                else:
                    prim[sku] = prim.get(sku, 0.0) + cant * mult * f

        acumular(cod, factor, ())
        masa = _base(unidad)[0] == "masa"
        return Resultado(cod, cantidad, unidad, cu * cantidad, cu, cu * (1000.0 if masa else 1.0), prim, falt, factor)

    # ------------------------------------------------------------------ nombres (contingencia)
    @staticmethod
    def norm_receta(t) -> str:
        t = unicodedata.normalize("NFKD", str(t or ""))
        t = "".join(c for c in t if not unicodedata.combining(c)).upper()
        t = re.sub(r"\(.*?\)", " ", t)
        t = re.sub(r"\b(GR|GRS|UN|UNIDAD|GRAMO|PORC|MADRE|PROD|OBRADOR|LIMPIA|KG|NUEVA|NUEVO)\b", " ", t)
        t = re.sub(r"[^A-Z0-9 ]+", " ", t)
        return " ".join(t.split())

    @staticmethod
    def contingencia(nombre: str) -> str:
        """Código de contingencia determinístico (mismo nombre -> mismo código): TEMP_<NOMBRE>_<hash4>."""
        k = MotorRecetas.norm_receta(nombre) or str(nombre or "").strip().upper()
        slug = re.sub(r"[^A-Z0-9]+", "_", k).strip("_")[:28]
        return f"TEMP_{slug}_{hashlib.sha1(k.encode('utf8')).hexdigest()[:4]}"

    @staticmethod
    def _tokens_cubiertos(q: list[str], c: list[str]) -> bool:
        """Cada token significativo de la consulta debe tener un par en el candidato (igual, mismo prefijo
        de 5+ letras, o ratio >= .85): 'ENCURTIDOS' ~ 'ENCURTIDA', pero 'BAKLAVA' !~ 'HALVA'."""
        stop = {"DE", "DEL", "LA", "EL", "Y", "CON", "A", "AL", "EN"}
        qs = [t for t in q if t not in stop]
        return bool(qs) and all(any(t == u or (len(t) >= 5 and len(u) >= 5 and t[:5] == u[:5]) or
                                    difflib.SequenceMatcher(None, t, u).ratio() >= 0.85 for u in c) for t in qs)

    def buscar_por_nombre(self, nombre: str, minimo: float = 0.85):
        """Aproxima un nombre contra NOM_RECETA/NOMBRE_RECETA. -> (cod_receta|None, similitud, alternativas).
        Acepta solo si similitud >= minimo Y todos los tokens significativos están cubiertos."""
        if self._idx_nombres is None:
            self._idx_nombres = {}
            for cod, rc in self.recetas.items():
                for nm in (rc["nombre"], rc["nom"]):
                    k = self.norm_receta(nm)
                    if k:
                        self._idx_nombres.setdefault(k, []).append(cod)
        k = self.norm_receta(nombre)
        if not k:
            return None, 0.0, []
        if len(k.split()) == 1:      # una sola palabra: solo coincidencia casi exacta (Hinojo != MADRE HINOJOS)
            minimo = max(minimo, 0.97)
        sc = sorted(((difflib.SequenceMatcher(None, k, c).ratio(), c) for c in self._idx_nombres), reverse=True)[:5]
        alts = [(self._idx_nombres[c][0], round(r, 2)) for r, c in sc if r >= 0.6]
        for r, c in sc:
            if r >= minimo and self._tokens_cubiertos(k.split(), c.split()):
                return self._idx_nombres[c][0], round(r, 3), alts
        return None, (round(sc[0][0], 3) if sc else 0.0), alts

    # ------------------------------------------------------------------ C_UN: fuente única de verdad
    def c_un_recalculado(self) -> dict[str, float]:
        """COD_RECETA -> C_UN (costo por unidad de rendimiento) recalculado por el BOM. Solo recetas completas."""
        out, memo = {}, {}
        for cod in self.recetas:
            cu, falt = self._costo_unit(cod, (), memo)
            if cu is not None and not falt:
                out[cod] = cu
        return out

    def sincronizar_c_un(self, src: Path, dst: Path, log_csv: Path | None = None) -> dict:
        """Escribe en `dst` (copia de Recetas.xlsx) el C_UN del BOM en Recetas!L (todas las filas de cada
        receta) y Costos!D. Parcheo de XML: el resto del libro queda intacto. -> estadísticas."""
        nuevo = self.c_un_recalculado()
        recetas_L: dict[int, float] = {}
        costos_D: dict[int, float] = {}
        filas_log = []
        for cod, cu in nuevo.items():
            rc = self.recetas[cod]
            viejo = rc["c_un_hoja"]
            filas_log.append([cod, rc["nombre"], viejo, cu, (cu / viejo - 1) if viejo else ""])
            for fila in self.filas_recetas.get(cod, []):
                recetas_L[fila] = cu
            if cod in self.fila_costos:
                costos_D[self.fila_costos[cod]] = cu
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            wbx = zin.read("xl/workbook.xml").decode("utf8")
            rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf8")

            def ruta(nombre):
                rid = re.search(r'<sheet\b[^>]*name="%s"[^>]*r:id="([^"]+)"' % nombre, wbx).group(1)
                t = (re.search(r'<Relationship\b[^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or
                     re.search(r'<Relationship\b[^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)).group(1).lstrip("/")
                return t if t.startswith("xl/") else "xl/" + t
            objetivo = {ruta("Recetas"): (get_column_letter(self.col_recetas_cun + 1), recetas_L),
                        ruta("Costos"): (get_column_letter(self.col_costos_cun + 1), costos_D)}
            parcheadas = {}
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename in objetivo:
                    col, vals = objetivo[item.filename]
                    xml = data.decode("utf8")
                    hechas = [0]

                    def fn(m, col=col, vals=vals, hechas=hechas):
                        fila = int(m.group(1))
                        if fila not in vals:
                            return m.group(0)
                        hechas[0] += 1
                        attrs = re.sub(r'\s+t="[^"]*"', "", m.group(2))
                        return f'<c r="{col}{fila}"{attrs}><v>{vals[fila]!r}</v></c>'
                    xml = re.sub(r'<c r="%s(\d+)"([^>/]*?)(?:/>|>.*?</c>)' % col, fn, xml, flags=re.S)
                    parcheadas[item.filename] = hechas[0]
                    data = xml.encode("utf8")
                elif item.filename == "xl/workbook.xml" and "fullCalcOnLoad" not in wbx:
                    data = wbx.replace("<calcPr", '<calcPr fullCalcOnLoad="1"', 1).encode("utf8")
                zout.writestr(item, data)
        if log_csv:
            with open(log_csv, "w", newline="", encoding="utf-8-sig") as fh:
                w = csv.writer(fh)
                w.writerow(["cod_receta", "nombre", "c_un_hoja", "c_un_motor", "delta_rel"])
                w.writerows(sorted(filas_log, key=lambda r: -abs(r[4]) if r[4] != "" else 0))
        difieren = sum(1 for r in filas_log if r[4] != "" and abs(r[4]) > 0.02)
        return dict(recetas=len(nuevo), celdas=parcheadas, difieren_2pct=difieren, total_recetas=len(self.recetas))


if __name__ == "__main__":  # auto-test: el costo recalculado debe coincidir con C_UN de la hoja
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    base = Path(__file__).resolve().parent / "Copias_Gestion" / "originales_export" / "Recetas.xlsx"
    m = MotorRecetas(base)
    ok = bad = inc = 0
    peores = []
    for cod, rc in m.recetas.items():
        ref = rc["c_un_hoja"]
        r = m.explotar(cod, rc["rend"] or 1)
        if not r or not ref:
            continue
        if not r.completo:
            inc += 1
            continue
        d = abs(r.costo_unitario / ref - 1) if ref else 0
        if d < 0.02:
            ok += 1
        else:
            bad += 1
            peores.append((d, cod, ref, r.costo_unitario))
    print(f"recetas={len(m.recetas)} insumos={len(m.insumos)} | coinciden con C_UN (±2%): {ok} | difieren: {bad} | incompletas: {inc}")
    for d, cod, ref, cu in sorted(peores, reverse=True)[:8]:
        print(f"   {cod}: hoja={ref:.4f} motor={cu:.4f} ({d:.0%})")
    r = m.explotar("M1", 2000)
    print("M1 x2000 gr: factor", r.factor, "costo", round(r.costo_total, 2), "precio KG", round(r.precio_gestion, 2),
          "| primarios:", {k: round(v, 1) for k, v in list(r.primarios.items())[:6]})
